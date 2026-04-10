from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, range_boundaries

from .models import LoadedWorkbook, SheetSpec, ThemeSpec
from .project import load_theme, load_workbook
from .utils import is_cell_address, is_range_address, sheet_refs_from_formula


def issue(
    severity: str,
    code: str,
    path: Path,
    message: str,
    field: str | None = None,
) -> dict[str, Any]:
    payload = {
        "severity": severity,
        "code": code,
        "path": str(path),
        "message": message,
    }
    if field:
        payload["field"] = field
    return payload


def _validate_style_name(
    issues: list[dict[str, Any]],
    style_name: str | None,
    theme: ThemeSpec | None,
    path: Path,
    field: str,
) -> None:
    if style_name and (theme is None or style_name not in theme.styles):
        theme_name = theme.name if theme else "default"
        issues.append(issue("error", "unknown_style", path, f"Style '{style_name}' is not defined in theme '{theme_name}'", field))


def validate_sheet_spec(
    sheet: SheetSpec,
    path: Path,
    theme: ThemeSpec | None = None,
    workbook_sheet_titles: set[str] | None = None,
    project_root: Path | None = None,
) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    workbook_sheet_titles = workbook_sheet_titles or set()

    for column, width in sheet.column_widths.items():
        if not column.isalpha() or not column.isupper():
            issues.append(issue("error", "invalid_column_reference", path, f"Column width key '{column}' is invalid", "/column_widths"))
        elif column_index_from_string(column) > 16384 or width <= 0:
            issues.append(issue("error", "invalid_column_reference", path, f"Column width entry '{column}' is out of range", "/column_widths"))

    for row, height in sheet.row_heights.items():
        try:
            row_number = int(row)
        except ValueError:
            row_number = 0
        if row_number < 1 or row_number > 1048576 or height <= 0:
            issues.append(issue("error", "invalid_row_reference", path, f"Row height entry '{row}' is out of range", "/row_heights"))

    for index, cell in enumerate(sheet.cells):
        if not is_cell_address(cell.cell):
            issues.append(issue("error", "invalid_cell_address", path, f"Cell address '{cell.cell}' is invalid", f"/cells/{index}/cell"))
        if cell.formula is not None and not cell.formula.startswith("="):
            issues.append(issue("error", "invalid_formula", path, "Formula must start with '='", f"/cells/{index}/formula"))
        if cell.formula:
            for ref in sheet_refs_from_formula(cell.formula):
                if ref not in workbook_sheet_titles:
                    issues.append(issue("warning", "unresolved_sheet_ref", path, f"Formula references sheet '{ref}' which is not in this workbook", f"/cells/{index}/formula"))
        _validate_style_name(issues, cell.style, theme, path, f"/cells/{index}/style")
        if cell.image_path and not Path(cell.image_path).is_absolute():
            candidate = (project_root / cell.image_path) if project_root else (path.parent / cell.image_path)
            if not candidate.exists():
                issues.append(issue("error", "missing_asset", path, f"Missing asset '{cell.image_path}'", f"/cells/{index}/image_path"))

    for index, range_spec in enumerate(sheet.ranges):
        if not is_cell_address(range_spec.anchor):
            issues.append(issue("error", "invalid_cell_address", path, f"Range anchor '{range_spec.anchor}' is invalid", f"/ranges/{index}/anchor"))
        for row_index, style in range_spec.row_styles.items():
            _validate_style_name(issues, style, theme, path, f"/ranges/{index}/row_styles/{row_index}")
        for col_index, style in range_spec.col_styles.items():
            _validate_style_name(issues, style, theme, path, f"/ranges/{index}/col_styles/{col_index}")

    for index, merge in enumerate(sheet.merges):
        if not is_range_address(merge):
            issues.append(issue("error", "invalid_range_address", path, f"Merge range '{merge}' is invalid", f"/merges/{index}"))

    table_ranges = []
    for index, table in enumerate(sheet.tables):
        if not is_range_address(table.ref):
            issues.append(issue("error", "invalid_range_address", path, f"Table range '{table.ref}' is invalid", f"/tables/{index}/ref"))
        else:
            table_ranges.append((table.ref, index))

    for merge_index, merge in enumerate(sheet.merges):
        if not is_range_address(merge):
            continue
        merge_bounds = range_boundaries(merge)
        for table_ref, table_index in table_ranges:
            table_bounds = range_boundaries(table_ref)
            if not (merge_bounds[2] < table_bounds[0] or merge_bounds[0] > table_bounds[2] or merge_bounds[3] < table_bounds[1] or merge_bounds[1] > table_bounds[3]):
                issues.append(issue("error", "merge_overlaps_table", path, f"Merge range '{merge}' overlaps table range '{table_ref}'", f"/merges/{merge_index}"))

    for index, chart in enumerate(sheet.charts):
        for series_index, series in enumerate(chart.series):
            for field_name in ["values", "categories"]:
                ref = getattr(series, field_name)
                if not ref:
                    continue
                if "!" not in ref:
                    continue
                sheet_name = ref.split("!", 1)[0].strip("'")
                if sheet_name not in workbook_sheet_titles:
                    issues.append(issue("warning", "unknown_chart_sheet_ref", path, f"Chart series references unknown sheet '{sheet_name}'", f"/charts/{index}/series/{series_index}/{field_name}"))
    return issues


def validate_workbook(loaded: LoadedWorkbook) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    seen_sheet_ids: set[str] = set()
    seen_table_names: set[str] = set()
    titles = {sheet.spec.title for sheet in loaded.sheets}

    if loaded.spec.theme and loaded.theme is None:
        issues.append(issue("error", "missing_theme", loaded.path, f"Theme '{loaded.spec.theme}' could not be resolved", "/theme"))

    for idx, sheet in enumerate(loaded.sheets):
        if sheet.spec.sheet_id in seen_sheet_ids:
            issues.append(issue("error", "duplicate_sheet_id", sheet.path, f"Duplicate sheet_id '{sheet.spec.sheet_id}'", "/sheet_id"))
        seen_sheet_ids.add(sheet.spec.sheet_id)
        issues.extend(validate_sheet_spec(sheet.spec, sheet.path, theme=loaded.theme, workbook_sheet_titles=titles, project_root=loaded.project_root))
        for table_index, table in enumerate(sheet.spec.tables):
            if table.name in seen_table_names:
                issues.append(issue("error", "duplicate_table_name", sheet.path, f"Duplicate table name '{table.name}'", f"/tables/{table_index}/name"))
            seen_table_names.add(table.name)

    return {"ok": not any(item["severity"] == "error" for item in issues), "issues": issues}


def validate_path(target_path: Path, project_root: Path | None = None) -> dict[str, Any]:
    target_path = target_path.resolve()
    if target_path.name == "workbook.json":
        loaded = load_workbook(target_path, project_root=project_root)
        return validate_workbook(loaded)

    sheet = SheetSpec.model_validate_json(target_path.read_text(encoding="utf-8"))
    theme = None
    if project_root:
        theme = load_theme(project_root, None)
    issues = validate_sheet_spec(sheet, target_path, theme=theme)
    return {"ok": not any(item["severity"] == "error" for item in issues), "issues": issues}
